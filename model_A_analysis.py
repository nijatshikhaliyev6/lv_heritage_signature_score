import openpyxl
import numpy as np
import pandas as pd
import statsmodels.api as sm

wb = openpyxl.load_workbook('LV_Heritage_Signature_Score__4_.xlsx', data_only=True)
ws = wb['3_OBJECTS']

rows = []
for r in range(5, 465):
    if ws.cell(r, 9).value != 'yes':
        continue
    rows.append(dict(
        obj_id=ws.cell(r, 1).value,
        name=ws.cell(r, 2).value,
        year=ws.cell(r, 3).value,
        iconic=ws.cell(r, 8).value,
        type=ws.cell(r, 45).value,
        S8=ws.cell(r, 43).value,
        S6=ws.cell(r, 46).value,
    ))

df = pd.DataFrame(rows)
df = df.dropna(subset=['year', 'S6', 'type', 'iconic'])
df['iconic_bin'] = (df['iconic'] == 'yes').astype(int)
df['year_c'] = (df['year'] - 1854) / 10  # decades since founding, for an interpretable intercept

# fold the 4 "other" objects into the modal category (trunk) to avoid a near-empty dummy
df['type_clean'] = df['type'].replace({'other': 'trunk'})

print('n =', len(df))
print(df['type_clean'].value_counts())
print(df['iconic_bin'].value_counts())
print(f"year range: {df['year'].min():.0f}-{df['year'].max():.0f}")

# ---------------- MODEL A (renamed): S6 ~ type + iconic + year, all 400 objects ----------------
exog = pd.get_dummies(df[['year_c', 'type_clean', 'iconic_bin']], columns=['type_clean'], drop_first=True).astype(float)
exog = sm.add_constant(exog)
exog = exog[['const', 'year_c', 'type_clean_case', 'type_clean_suitcase', 'type_clean_trunk', 'iconic_bin']]

endog = df['S6'].astype(float)

model = sm.OLS(endog, exog)
res = model.fit(cov_type='HC3')  # heteroskedasticity-robust SEs, standard practice for this sample size

print(res.summary())

table = pd.DataFrame({
    'coef': res.params, 'se': res.bse, 't': res.tvalues, 'p': res.pvalues,
    'ci_lo': res.conf_int()[0], 'ci_hi': res.conf_int()[1]
})
table.to_csv('modelA_full400.csv')
df.to_csv('modelA_full400_data.csv', index=False)
print('\nSaved modelA_full400.csv, modelA_full400_data.csv')
print(f'\nR-squared = {res.rsquared:.3f}, Adj. R-squared = {res.rsquared_adj:.3f}, n = {res.nobs:.0f}')
